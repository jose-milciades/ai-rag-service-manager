import io

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.infrastructure.ocr.vision_ocr import PAGE_BREAK
from app.services.embedding.text_extractors import (
    NeedsOcrError,
    UnsupportedFormatError,
    extract_text,
    needs_ocr,
)


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Politica de privacidad de prueba")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "clave"
    table.rows[0].cells[1].text = "valor"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["producto", "precio"])
    ws.append(["lapiz", 10])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Titulo de diapositiva"
    slide.placeholders[1].text = "Contenido de la diapositiva"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_docx_extracts_paragraphs_and_tables():
    text = extract_text(_docx_bytes(), "a.docx")
    assert "Politica de privacidad de prueba" in text
    assert "clave | valor" in text


def test_xlsx_extracts_rows():
    text = extract_text(_xlsx_bytes(), "a.xlsx")
    assert "Hoja: Datos" in text and "lapiz | 10" in text


def test_pptx_extracts_slide_text():
    text = extract_text(_pptx_bytes(), "a.pptx")
    assert "Titulo de diapositiva" in text and PAGE_BREAK not in text


def test_extraction_is_deterministic():
    data = _docx_bytes()
    assert extract_text(data, "a.docx") == extract_text(data, "a.docx")


def test_unknown_extension_fails_explicitly():
    with pytest.raises(UnsupportedFormatError):
        extract_text(b"\x00\x01", "a.bin")


def test_image_without_ocr_fails_explicitly():
    with pytest.raises(NeedsOcrError):
        extract_text(b"\x89PNG", "a.png", ocr=None)


def test_text_too_large_is_rejected():
    with pytest.raises(ValueError, match="demasiado grande"):
        extract_text(b"a" * 100, "a.txt", max_chars=10)


def test_needs_ocr_only_for_pdf_and_images():
    assert needs_ocr("a.PDF") and needs_ocr("a.jpg")
    assert not needs_ocr("a.docx")


def test_pptx_soft_line_break_is_normalized():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Linea uno\x0bLinea dos"
    buf = io.BytesIO()
    prs.save(buf)
    text = extract_text(buf.getvalue(), "a.pptx")
    assert "\x0b" not in text and "Linea uno" in text and "Linea dos" in text
