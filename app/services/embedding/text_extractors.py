"""Extraccion de texto por formato para indexacion.

Registro ``extension -> funcion``. Es la unica fuente de texto tanto para
indexar como para reconstruir contexto (``expand_context`` usa los offsets
``start_index/end_index`` guardados al indexar), por lo que debe ser
**determinista**: el mismo archivo produce siempre el mismo texto. Por eso el
resultado de un OCR se cachea en storage (ver ``DocumentEmbeddingService``).

Reemplaza a Spire (licencia de pago) de ``edi-ai-analysis-ai`` por librerias
abiertas. Una extension desconocida falla con ``UnsupportedFormatError`` en vez
de decodificar bytes binarios como texto.
"""

import io
import logging
import re
import shutil
import subprocess  # nosec B404 - catdoc/catppt con lista de args, sin shell
import tempfile
from pathlib import Path

import pdfplumber

from app.infrastructure.ocr.vision_ocr import PAGE_BREAK, VisionOcr

logger = logging.getLogger(__name__)

_PLAIN_TEXT = {"txt", "md", "json", "csv", "py", "yaml", "yml", "html", "xml"}
_IMAGES = {"png", "jpg", "jpeg", "tif", "tiff", "bmp", "gif", "webp"}
_MIN_READABLE_RATIO = 0.5
_SUBPROCESS_TIMEOUT = 120


class UnsupportedFormatError(ValueError):
    """Extension sin extractor: no se indexa en silencio."""


class NeedsOcrError(ValueError):
    """El documento no tiene texto extraible y el OCR no esta disponible."""


def extract_text(
    content: bytes,
    file_name: str,
    ocr: VisionOcr | None = None,
    min_chars_per_page: int = 20,
    max_chars: int | None = None,
) -> str:
    """Extrae texto del archivo segun su extension. ``ocr`` solo se usa si hace falta."""
    ext = _resolve_extension(content, file_name)
    if ext in _PLAIN_TEXT:
        text = content.decode("utf-8", errors="ignore")
    elif ext == "pdf":
        text = _pdf(content, ocr, min_chars_per_page)
    elif ext in _IMAGES:
        text = _image(content, ocr)
    elif ext == "docx":
        text = _docx(content)
    elif ext == "xlsx":
        text = _xlsx(content)
    elif ext == "pptx":
        text = _pptx(content)
    elif ext == "xls":
        text = _xls(content)
    elif ext == "doc":
        text = _legacy_office(content, "catdoc", ".doc")
    elif ext == "ppt":
        text = _legacy_office(content, "catppt", ".ppt")
    else:
        raise UnsupportedFormatError(f"Formato no soportado para indexar: .{ext or '?'}")

    if max_chars is not None and len(text) > max_chars:
        raise ValueError(f"Texto extraido demasiado grande ({len(text)} > {max_chars} caracteres)")
    if text.strip() and ext not in _PLAIN_TEXT and _readable_ratio(text) < _MIN_READABLE_RATIO:
        raise ValueError("El texto extraido no es legible (posible binario); no se indexa")
    return text


def needs_ocr(file_name: str, content: bytes | None = None) -> bool:
    """Formatos cuyo resultado puede depender de OCR (y por eso se cachea)."""
    ext = _resolve_extension(content, file_name)
    return ext == "pdf" or ext in _IMAGES


def _resolve_extension(content: bytes | None, file_name: str) -> str:
    """Extension efectiva: manda el contenido (firma) sobre el nombre.

    Hay archivos con extension errada (un PDF llamado ``.pptx``) o sin
    extension real (``Informe_V1.7``). ``analysis-ai`` los resolvia con
    ``magic``; aqui se detectan por firma los formatos que se soportan.
    """
    by_name = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    if not content:
        return by_name
    detected = _sniff(content)
    if detected is None or detected == by_name:
        return by_name
    # Un .docx/.xlsx/.pptx es un ZIP: solo se corrige si el nombre no es de OOXML.
    if detected == "zip":
        return by_name if by_name in {"docx", "xlsx", "pptx"} else _sniff_ooxml(content) or by_name
    # OLE sirve para doc/xls/ppt: no se distingue por firma, se respeta el nombre.
    if detected == "ole":
        return by_name
    logger.info("extension corregida por contenido: .%s -> .%s", by_name, detected)
    return detected


def _sniff(content: bytes) -> str | None:
    head = content[:8]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"
    if head.startswith(b"\x89PNG"):
        return "png"
    if head.startswith(b"\xff\xd8\xff"):
        return "jpg"
    return None


def _sniff_ooxml(content: bytes) -> str | None:
    import zipfile

    try:
        names = zipfile.ZipFile(io.BytesIO(content)).namelist()
    except zipfile.BadZipFile:
        return None
    for prefix, ext in (("word/", "docx"), ("xl/", "xlsx"), ("ppt/", "pptx")):
        if any(n.startswith(prefix) for n in names):
            return ext
    return None


def _readable_ratio(text: str) -> float:
    sample = text[:20000]
    ok = sum(1 for ch in sample if ch.isalnum() or ch.isspace() or ch in ".,;:()-_/%$@#&+*'\"|")
    return ok / max(1, len(sample))


def _pdf(content: bytes, ocr: VisionOcr | None, min_chars: int) -> str:
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        pages = [(page.extract_text() or "") for page in pdf.pages]
    empty = [i + 1 for i, t in enumerate(pages) if len(t.strip()) < min_chars]
    if empty:
        if ocr is None or not ocr.enabled:
            if len(empty) == len(pages):
                raise NeedsOcrError("PDF sin texto extraible (escaneado) y OCR no disponible")
        else:
            for number, text in ocr.ocr_pdf_pages(content, empty).items():
                pages[number - 1] = text
    return PAGE_BREAK.join(pages)


def _image(content: bytes, ocr: VisionOcr | None) -> str:
    if ocr is None or not ocr.enabled:
        raise NeedsOcrError("Imagen sin OCR disponible")
    return ocr.ocr_image(content)


def _docx(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = [f"Hoja: {ws.title}"]
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        sheets.append("\n".join(rows))
    wb.close()
    return PAGE_BREAK.join(sheets)


def _xls(content: bytes) -> str:
    import xlrd

    wb = xlrd.open_workbook(file_contents=content)
    sheets = []
    for ws in wb.sheets():
        rows = [f"Hoja: {ws.name}"]
        for r in range(ws.nrows):
            cells = [str(c).strip() for c in ws.row_values(r)]
            if any(cells):
                rows.append(" | ".join(cells))
        sheets.append("\n".join(rows))
    return PAGE_BREAK.join(sheets)


def _pptx(content: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(content))
    slides = []
    for slide in prs.slides:
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                # PowerPoint usa \x0b (salto vertical) como salto de linea suave.
                parts.append(shape.text_frame.text.replace("\x0b", "\n"))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append(slide.notes_slide.notes_text_frame.text)
        slides.append("\n".join(parts))
    return PAGE_BREAK.join(slides)


def _legacy_office(content: bytes, tool: str, suffix: str) -> str:
    """``.doc``/``.ppt`` binarios via ``catdoc``/``catppt`` (paquete Debian ``catdoc``)."""
    binary = shutil.which(tool)
    if binary is None:
        raise UnsupportedFormatError(f"Falta el binario '{tool}' para leer {suffix}")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / f"input{suffix}"
        path.write_bytes(content)
        result = subprocess.run(  # nosec B603 - binario resuelto con which, args fijos
            [binary, "-d", "utf-8", str(path)],
            capture_output=True,
            timeout=_SUBPROCESS_TIMEOUT,
            check=False,
        )
    if result.returncode != 0:
        raise ValueError(f"{tool} fallo: {result.stderr.decode('utf-8', 'ignore')[:200]}")
    return re.sub(r"\n{3,}", "\n\n", result.stdout.decode("utf-8", errors="ignore"))
